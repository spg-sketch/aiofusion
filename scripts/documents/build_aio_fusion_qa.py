"""Build separate sales and internal-review Word drafts from approved editorial inputs.

The original uploads are read only. The sales file is a new package, never a
redacted copy of the original Word document.
"""
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "deliverables/aio-fusion-qa-v1.5"
SOURCE = ROOT / "attached_assets/AIO-Fusion-Master-QA-sales_conversations_1789740507981.docx"
EMAIL = ROOT / "attached_assets/Pasted-Drafting-that-email-now-you-ll-be-able-to-paste-it-as-i_1789740530597.txt"
DATA = json.loads((ROOT / "docs/editorial/aio-fusion-qa-v1.5.json").read_text())
NAVY = "18384B"
TEAL = "157B83"
GREY = "52616B"
DATE = "18 September 2026"
HASHES = {
    SOURCE.name: "617d0a6ef572ba67950d372d23ff1bc3f0157a098c15bd74cd2d83ba063e65a6",
    EMAIL.name: "1ae0268bd97aaa7a9978cba363f6852ce698461a14baf53397aed7c6d1c6400f",
}


def field(paragraph, instruction):
    r = paragraph.add_run()
    f = OxmlElement("w:fldSimple")
    f.set(qn("w:instr"), instruction)
    r._r.addnext(f)


def document(internal=False):
    d = Document()
    # python-docx's stock template includes a bibliography custom XML part.
    # Neither deliverable needs it, so omit it from the new package entirely.
    for rid, rel in list(d.part.rels.items()):
        if rel.reltype.endswith("/customXml"):
            d.part.drop_rel(rid)
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)
    s.top_margin = s.bottom_margin = Inches(0.72)
    s.left_margin = s.right_margin = Inches(0.78)
    s.header_distance = s.footer_distance = Inches(0.32)
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10 if internal else 10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.widow_control = True
    for name, size in [("Title", 28), ("Subtitle", 13), ("Heading 1", 20), ("Heading 2", 13), ("Heading 3", 11)]:
        st = d.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(NAVY)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
    d.styles["Heading 3"].font.color.rgb = RGBColor.from_string(TEAL)
    st = d.styles.add_style("Source quotation", 1)
    st.font.name = "Calibri"
    st.font.size = Pt(9.5)
    st.font.color.rgb = RGBColor.from_string(GREY)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.widow_control = True
    h = s.header.paragraphs[0]
    h.text = "AIO FUSION  |  INTERNAL REVIEW - DO NOT SEND TO PROSPECTS" if internal else "AIO FUSION  |  MASTER Q&A v1.5"
    h.runs[0].font.size = Pt(8)
    h.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    f = s.footer.paragraphs[0]
    f.text = "INTERNAL ONLY - separate from prospect copy" if internal else "DRAFT FOR OWNER APPROVAL - NOT APPROVED FOR RELEASE"
    for r in f.runs:
        r.font.size = Pt(8)
    f.add_run("  |  ").font.size = Pt(8)
    field(f, "PAGE")
    cp = d.core_properties
    cp.title = "AIO Fusion Master Q&A v1.5" + (" - Internal annotated review" if internal else "")
    cp.subject = "Internal editorial review" if internal else "Sales-ready copy: draft for owner approval"
    cp.author = "AIO Fusion"
    cp.last_modified_by = "AIO Fusion"
    cp.comments = ""
    cp.keywords = ""
    cp.category = "Internal" if internal else "Owner approval draft"
    cp.created = cp.modified = datetime(2026, 9, 18, tzinfo=timezone.utc)
    cp.revision = 1
    return d


def para(d, text, style=None):
    assert "\u2014" not in text, "New copy must not contain em dashes"
    return d.add_paragraph(text, style)


def table(d, headers, rows, source=False):
    t = d.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    width = 6.71 / len(headers)
    for c in t.columns:
        c.width = Inches(width)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), NAVY)
        cell._tc.get_or_add_tcPr().append(shade)
        for r in cell.paragraphs[0].runs:
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        for cell, text in zip(t.add_row().cells, values):
            if not source:
                assert "\u2014" not in text
            cell.text = text
    for row in t.rows:
        prop = row._tr.get_or_add_trPr()
        prop.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.space_before = Pt(5)
                for r in p.runs:
                    r.font.size = Pt(9.5)
    d.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def answer(d, item):
    for text in item["answer"]:
        para(d, text)
    if item.get("table"):
        table(d, **item["table"])
    for text in item.get("after", []):
        para(d, text)


def new_page_heading(d, text):
    d.add_page_break()
    d.add_heading(text, 1)


def clean_document():
    d = document()
    d.add_heading("AIO Fusion\nMaster Q&A v1.5", 0)
    para(d, "Sales-ready copy | Draft for owner approval", "Subtitle")
    para(d, DATE)
    para(d, "Status: not approved for release. This editorial draft is intended for sales conversations and prospect follow-up after owner approval. It is not a binding offer, a security certification or a guarantee of results.")
    para(d, "Commercial terms and dated beta arrangements are preserved from the supplied v1.4 and remain subject to confirmation. Where a factual answer is not yet established, the draft says so rather than filling the gap.")
    d.add_heading("How to read the answers", 2)
    para(d, "An observation is something found in the sampled AI response or fetched website content. An AI assessment interprets evidence and context. A planning estimate uses assumptions to help organise future work. These are different forms of information and should not be treated as interchangeable.")
    d.add_heading("Contents", 2)
    for s in DATA["sections"]:
        para(d, f'{s["number"]}. {s["title"]}')
    for s in DATA["sections"]:
        new_page_heading(d, f'{s["number"]}  {s["title"]}')
        for i, item in enumerate(s["items"], 1):
            d.add_heading(f'{s["number"]}.{i}  {item["question"]}', 2)
            answer(d, item)
    return d


CHECKLIST = [
    ("Measurement owner | prompt independence and identity",
     "Resolve section 2's 'roughly 99% blind' and no-reference wording against section 3's identity anchors and project-derived question generation. On the source's maximum described mix, eight category questions plus one identity question gives 8/9, about 88.9%, category questions, not 99%, if counted equally. Four responses per question would give 32/36, also about 88.9%. This arithmetic is only a contradiction check, not verification of the implemented method or of blindness. Examine actual prompt payloads and identify which stages receive names, domains, sectors and messages."),
    ("Measurement owner | aggregation, sampling and final assessment",
     "Verify actual run counts, question weights, scoring provider, dimension treatment, denominators and handling of failures, duplicates and missing answers. Reconcile 'averages' with the source's more specific aggregation description. Distinguish question weighting from final index construction. Document methodology versions and what prevents invalid comparisons."),
    ("Measurement owner | repeatability and API scope",
     "Obtain repeat-run evidence before saying robust, stable or repeatable. Temperature zero is not a guarantee of deterministic output, and API defaults are not proof of consumer-chat equivalence. Confirm actual provider/model identifiers, settings and retrieval conditions. Verify or discard the named model variants in the source; do not quietly replace them with guessed names."),
    ("Product owner | Planner and content scores",
     "Resolve model-generated versus configured Planner scoring. The clean draft follows the source's more specific description of activity, channels, workflow status and configurable values, pending confirmation. Confirm whether publication-time content scoring exists and keep it separate from projected programme value, article quality and observed audit assessment."),
    ("Research/marketing owner | external substantiation",
     "Do not reinstate the 80–90% earned-media assertion, earned-media-first hierarchy, first/unique positioning, blanket competitor comparisons, model suitability by client type, social weighting or 'most searches' claim without dated evidence and defined scope. Generated questions are not actual customer-demand or search-volume data."),
    ("Product/reporting owner | customer-visible output",
     "Demonstrate current live report screens and exports: provider splits, citations, mention rates, source context, spokesperson/content/message breakdowns, and saved history. Verify any append-only or reconstruct-any-history promise and the scope of actual retention. Backend storage or staging-only code is not proof of a customer-visible live feature. Do not reinstate causal-attribution language."),
    ("Product/web audit owner | measured versus assessed coverage",
     "Confirm which pages and resources are fetched, which checks are measured, and what supports any page-speed statement. Verify downloadable briefs and Word exports and the real save/handoff behaviour. Distinguish parsed facts from AI-generated recommendations; do not imply whole-site coverage from a page fetch."),
    ("Data owner | media and market intelligence",
     "Confirm in-house provenance, UK/US and sector coverage, the source's 125-category total, listing horizon, contact accuracy evidence and scoring labels. Do not promise every category or record refreshes within 30 days without operational evidence. Check organiser links, dates and deadlines; relevance indicators are not measured GEO impact."),
    ("Commercial owner | prices and allowances",
     "Approve the exact annual and quarterly annual-total figures, included Premium projects, VAT position, additional-project prices and 50/75/150 monthly allowances. Confirm chargeable actions, exclusions, resets, repeated actions, limits, support, contract duration, extra charges and applicable terms. All source prices are retained, not validated against billing configuration."),
    ("Commercial owner | beta",
     "Confirm the two-month free trial, structured feedback requirement, 30% first-year reduction, eligibility, 5 October 2026 opening, December endpoint and early-October support. Reconcile calendar dates with individual trial start/end dates. Specify any payment, cancellation, renewal and conversion terms. Do not turn an expectation of continuation into a contractual obligation."),
    ("Product/commercial owners | availability and roadmap",
     "Confirm current module availability, white-label package inclusion and scope, scoring customisation, report branding, support routing and Ask GEOrge behaviour. Remove unsupported promises from the approved copy. Additional providers, a release gateway, LinkedIn distribution, videos and timing remain proposals, not current deliverables. No new roadmap commitment is created."),
    ("Security owner | deployed controls and processors",
     "Verify hosting/database/backup providers, all actual processors, locations, subprocessor chains, data flows, provider training/retention terms, cross-border transfers and agreements. Check access isolation, staff/admin access, password/session controls, credential handling and audit logs in the deployed service. Document encryption in transit and at rest by service; do not infer it from a provider name."),
    ("Privacy/legal owner | rights, retention and contact data",
     "Confirm controller/processor roles, privacy notice currency, lawful sourcing and outreach use, correction/objection/suppression routes, subject-access handling and response timing. Establish return/export, live/backups/logs retention, deletion limits, incident procedures and legal holds. A general privacy notice does not substantiate a GDPR-compliance absolute."),
    ("Leadership/editorial owner | final approval",
     "Confirm named leadership roles and any biographies to restore, then approve external wording and the offer. Check approved claims against current customer-facing functionality, not just staging. Keep all unverified matters explicitly qualified or remove them before release; record approver, evidence date and approved version.")
]


def original_blocks(d, src, start, end):
    """Copy visible text and tables only, never comments/metadata or original XML."""
    indexes = {p._p: i for i, p in enumerate(src.paragraphs)}
    current = -1
    for el in src.element.body:
        if el.tag == qn("w:p"):
            current = indexes[el]
            if start <= current < end:
                text = Paragraph(el, src).text.strip()
                if text:
                    d.add_paragraph(text, "Source quotation")
        elif el.tag == qn("w:tbl") and start <= current < end:
            rows = [[c.text for c in r.cells] for r in Table(el, src).rows]
            table(d, rows[0], rows[1:], source=True)


def internal_document(src):
    d = document(True)
    d.add_heading("AIO Fusion Master Q&A v1.5\nInternal annotated review", 0)
    para(d, "INTERNAL ONLY | Must not accompany the prospect copy", "Subtitle")
    para(d, DATE)
    para(d, "Purpose: constructive editorial comparison with Master Q&A v1.4, dated 17 September 2026, and the accompanying email draft. The original provides a useful integrated-workflow explanation, practical commercial detail and important limitations. This revision preserves that substance while separating routine sales answers from confidential mechanics, commercial modelling and detailed procurement evidence.")
    para(d, "Status: draft for owner approval, not approved for release. This is an editorial exercise, not a product, security or legal audit. No live capability, commercial configuration or third-party claim has been independently certified. Neither original upload has been changed. No documents have been published or sent externally.")
    para(d, "The comparison covers 51 source questions or content blocks across all ten sections, including the boilerplate, explanatory subsections and unheaded Market Intelligence paragraph. Full visible source wording is reproduced by block, followed by the exact v1.5 replacement and rationale. Source quotations retain original punctuation, including any original em dashes; new copy does not use em dashes.")
    d.add_heading("Audience and handling", 1)
    table(d, ["Audience", "Appropriate material", "Handling"], [
        ["Routine prospect conversations", "Owner-approved clean v1.5: workflow, scope, limitations, commercial terms.", "Send only the clean file after approval. Do not attach this review."],
        ["Procurement, security and privacy reviewers", "Accurate hosting, processor, location, access, retention and privacy evidence.", "Answer legitimate questions with verified facts. Use proportionate confidentiality, not blanket secrecy."],
        ["Authorised technical evaluators", "Only verified methodology detail necessary for a defined evaluation.", "Scope the question and approve the disclosure. An NDA does not automatically authorise the full methodology."],
        ["Internal commercial leadership", "Fee bands, tier sizing, expected workloads and pricing assumptions.", "Keep internal. These explain commercial decisions, not customer entitlements."]
    ])
    d.add_heading("How decisions are labelled", 2)
    para(d, "Retain: preserve useful substance. Simplify: remove excess detail or narrow the claim. Restrict to procurement: move detailed factual evidence into a relevant due-diligence response. Restrict to confidential methodology: keep implementation recipes under controlled review. Factual confirmation required: the source is inconsistent, unsupported, time-sensitive or not verified. Labels can overlap. Internal commercial modelling is excluded through simplification and held for commercial leadership, not mislabelled as procurement evidence.")
    d.add_heading("Front matter and approval status", 2)
    d.add_heading("Original wording", 3)
    d.add_paragraph(src.paragraphs[3].text, "Source quotation")
    d.add_paragraph(src.paragraphs[6].text, "Source quotation")
    d.add_heading("Replacement wording", 3)
    para(d, "Status: not approved for release. This editorial draft is intended for sales conversations and prospect follow-up after owner approval. It is not a binding offer, a security certification or a guarantee of results.")
    d.add_heading("Reason: simplify; factual confirmation required", 3)
    para(d, "The original internal-only label conflicts with its near-verbatim reuse instruction. Build a genuinely separate prospect draft, keep approval visible on every page and do not rely on a confidentiality label to cure inappropriate sales claims.")
    new_page_heading(d, "Unresolved-claims checklist")
    para(d, "All items below are open until the relevant owner supplies evidence or approves narrower wording. This checklist is deliberately separate from confidentiality decisions: withholding detail does not make an unsupported assertion true.")
    for i, (title, body) in enumerate(CHECKLIST, 1):
        d.add_heading(f"{i}. {title}", 2)
        para(d, body)
    new_page_heading(d, "Withheld-material register")
    para(d, "Exact source wording is retained in the section-by-section comparison below. This register identifies who should review it and what must be verified before any selective disclosure. It is not a completed technical manual or security pack.")
    table(d, ["Source location and material", "Intended audience", "Confirmation needed"], [
        ["Sections 2–3: blind/branded ratio, identity anchors and question-generation inputs.", "Measurement owner; authorised technical evaluator only if necessary.", "Actual prompt payloads and exposure of company information at each stage. No unverified blindness claim."],
        ["Section 3: exact question weights, probe/run counts, temperature settings, scoring-provider implementation and agency-specific weights.", "Measurement owner; scoped confidential methodology discussion.", "Current configuration, aggregation, score construction, repeatability evidence and comparability effects."],
        ["Section 5: in-house agentic construction and rolling 30-day refresh process.", "Data owner; selective technical evaluator.", "Separate proprietary build detail from buyer-relevant provenance, lawful sourcing and evidenced currency."],
        ["Section 9: fee bands, workload examples, calculated action totals, cap headroom and expected additional-project volumes.", "Internal commercial leadership only.", "Internal assumptions need not be externally disclosed. Customer prices and allowances still require approval."],
        ["Section 10: database technology, hashing algorithm, cookie flags, session lifetime/revocation and IP handling.", "Security and procurement reviewers.", "Verify deployed controls. Do not withhold material facts required for risk assessment merely because details are technical."],
        ["Sections 2–3, 5, 8–9: model variants, provider plans, white-label/scoring promises, gateway dates and video plans.", "Product/commercial owners; approved public statements only.", "Actual availability, authorised scope and dates. Proposed features do not become promises through an NDA."]
    ])
    d.add_heading("Accompanying email: disposition of recommendations", 2)
    table(d, ["Email point", "Editorial response"], [
        ["1. Authority methodology", "Accepted the audience separation. Retained high-level dimensions and observation/assessment distinction. Exact mechanics remain unverified internal material, not an automatic full-methodology handout under NDA."],
        ["2. Blind-probe design", "Qualified rather than accepted. The email calls 99% blind a differentiator, but the source contradicts that claim. Also rejected the categorical assertion that setup never affects prompts until payloads are checked."],
        ["3. Media construction", "Removed the build recipe and 30-day guarantee. Kept source-attributed provenance and explicit record-currency limitations."],
        ["4. Tier modelling", "Kept all supplied price points and action caps for confirmation. Removed fee bands and workload calculations from the prospect draft."],
        ["5. Infrastructure/authentication", "Simplified configuration detail but retained named hosting/processors and privacy gaps. Exact hosting and controls remain legitimate procurement disclosures, not inherently secret."],
        ["6. Roadmap", "Removed speculative model names, dates and channel promises. Retained current coverage and explicit beta limitations; proposed features remain proposals."]
    ])
    for s in DATA["sections"]:
        new_page_heading(d, f'Comparison: {s["number"]}  {s["title"]}')
        for i, item in enumerate(s["items"], 1):
            d.add_heading(f'{s["number"]}.{i}  {item["question"]}', 2)
            para(d, "Decision: " + item["category"])
            d.add_heading("Original v1.4 wording", 3)
            original_blocks(d, src, item["start"], item["end"])
            d.add_heading("Replacement v1.5 wording", 3)
            answer(d, item)
            d.add_heading("Reason and handling", 3)
            para(d, item["reason"])
    new_page_heading(d, "Editorial approval and release checklist")
    for text in [
        "Confirm the open facts above with the named functional owners. Keep evidence and its date with the approval record.",
        "Resolve or explicitly qualify commercial terms, beta dates and module availability. Do not silently update source prices or substitute a staging feature for a live capability.",
        "Approve the clean wording for the intended audience. Check that observations, AI assessments and planning estimates remain distinct throughout.",
        "Check the final Word rendering, tables, section sequence and version labels after any owner edits.",
        "Inspect the final clean file for comments, tracked changes, hidden text, embedded objects and unintended document properties. Regenerate cleanly rather than hiding deleted text.",
        "Release only the separately approved clean file. Keep this internal review, source documents and withheld material out of prospect attachments.",
        "For an NDA discussion, record the buyer's question, necessity, approved scope, approver and evidence version. Disclose only what is needed, not the whole internal review."
    ]:
        para(d, "[ ] " + text)
    table(d, ["Approval record", "To be completed by owners"], [
        ["Commercial and beta terms", "Name / date / evidence / decision"],
        ["Product and measurement claims", "Name / date / evidence / decision"],
        ["Security and privacy claims", "Name / date / evidence / decision"],
        ["Final external release", "Approver / date / approved version / intended audience"]
    ])
    d.add_heading("Source record", 2)
    para(d, "Source A: AIO-Fusion-Master-QA-sales_conversations_1789740507981.docx, labelled v1.4, 17 September 2026.")
    para(d, "Source B: Pasted-Drafting-that-email-now-you-ll-be-able-to-paste-it-as-i_1789740530597.txt, accompanying editorial email draft.")
    para(d, "Scope excludes changes to the application, website, data, billing configuration and roadmap. This review does not provide legal certification, an IP-protection guarantee, an NDA or a complete procurement pack.")
    return d


def validate(clean_path, review_path, src):
    from lxml import etree
    report = []
    for p in [SOURCE, EMAIL]:
        assert hashlib.sha256(p.read_bytes()).hexdigest() == HASHES[p.name]
    report.append("Both original uploads match their pre-edit SHA-256 hashes.")
    items = [i for s in DATA["sections"] for i in s["items"]]
    covered = {n for i in items for n in range(i["start"], i["end"])}
    # Every body paragraph is classified; only headings, blanks and editorial preamble are exempt.
    exempt = {20, 49, 69, 71, 135, 173, 185, 188, 214, 225, 253}
    missing = [n for n, p in enumerate(src.paragraphs) if n >= 20 and p.text.strip() and n not in covered and n not in exempt]
    assert not missing, missing
    report.append(f"All ten sections and {len(items)} original questions/content blocks are mapped; no unclassified body text.")
    forbidden_elements = ["ins", "del", "moveFrom", "moveTo", "vanish", "webHidden", "commentRangeStart", "commentReference", "object", "altChunk"]
    for path in [clean_path, review_path]:
        with zipfile.ZipFile(path) as z:
            assert z.testzip() is None
            names = z.namelist()
            assert not any("comments" in n or "customXml" in n or "embeddings" in n for n in names)
            for n in names:
                if n.endswith(".xml"):
                    root = etree.fromstring(z.read(n))
                    assert not any(root.findall(".//" + qn("w:" + tag)) for tag in forbidden_elements), (path, n)
            docxml = z.read("word/document.xml").decode()
            if path == clean_path:
                assert "\u2014" not in docxml
                for restricted in ["99%", "1.5 times", "1.0 times", "0.5 times", "36 individual", "temperature", "identity anchor", "£4k", "£8k", "£16k", "31 actions", "62 actions", "134 actions", "scrypt", "HttpOnly", "SameSite", "30-day lifetime", "Sol, Terra", "Fable", "mid-2027", "agentic process", "30-day basis"]:
                    assert restricted.lower() not in docxml.lower(), restricted
    report.append("Both DOCX packages open and parse; no comments, tracked revisions, hidden-text flags, embedded objects or custom XML.")
    report.append("Clean package uses fresh properties and contains no source metadata or withheld recipes. No em dashes in clean text.")
    clean = Document(clean_path)
    assert len(clean.tables) == 3
    text = "\n".join(p.text for p in clean.paragraphs) + "\n".join(c.text for t in clean.tables for r in t.rows for c in r.cells)
    for value in ["£4,000", "£4,600", "£5,000", "£5,750", "£333", "£383", "£417", "£479", "£500", "£650", "£800", "50", "75", "150", "30%", "5 October 2026", "two-month", "December 2026"]:
        assert value in text, value
    report.append("All source plan/tier prices, monthly equivalents, allowances and beta financial/date terms retained with confirmation caveats.")
    (OUT / "verification.txt").write_text("\n".join(report) + "\nRendering check follows separately.\n")
    print("\n".join(report))


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    src = Document(SOURCE)
    clean_path = OUT / "AIO-Fusion-Master-QA-v1.5-OWNER-APPROVAL-DRAFT.docx"
    review_path = OUT / "AIO-Fusion-Master-QA-v1.5-INTERNAL-REVIEW-DO-NOT-SHARE.docx"
    clean_document().save(clean_path)
    internal_document(src).save(review_path)
    validate(clean_path, review_path, src)
    print(clean_path)
    print(review_path)